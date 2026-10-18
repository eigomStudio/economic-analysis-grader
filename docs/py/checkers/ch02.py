"""第2章 章末課題(世帯主の年齢階級別の消費支出: 平均・中央値・標準偏差・変動係数)のチェッカー。
grader.py と同じフォルダの checkers/ に置くと読み込まれる(Webアプリでも publish.py でコピーされる)。

評価の項目(S〜D、C以上が合格)
- 関数: 平均=AVERAGE、中央値=MEDIAN、標準偏差=STDEV.S(または STDEV)、変動係数=標準偏差/平均。値も照合する。
- 書式: 平均・中央値・標準偏差は桁区切り・小数第1位(#,##0.0)。変動係数は小数第1位または第2位。
- 考察: 字数と語句(変動係数・ばらつきへの言及、具体的な費目名)で判定する(LLMは使わない)。
  模範解答の「注意事項」のテキストボックス(授業での補足)は、考察に数えない。
"""
import re
import zipfile
import statistics
import unicodedata
from collections import Counter
from openpyxl import load_workbook

STATS = {"平均": "avg", "中央値": "med", "範囲": "rng", "標準偏差": "sd", "変動係数": "cv"}
FUNC = {"avg": r"AVERAGE\(", "med": r"MEDIAN\(", "rng": r"MAX\(.*MIN\(", "sd": r"(?:_XLFN\.)?STDEV(?:\.S)?\("}
MIN_ESSAY, MIN_ESSAY_NONE = 100, 10
ESSAY_DEDUCT = 50          # 考察がこの字数未満(明らかに短い)のときだけ減点する。それ以外の考察の不足(語句など)はコメントのみ
NEED_ITEMS = 2                    # 考察で触れてほしい費目の数
HAND_TYPED_MAX = 0                # 手入力の数値がこの個数を超えると「式なし」=D(第1章と同じ方針)
FMT_OK_RATIO = 0.9                # 書式が正しいセルがこの割合未満なら「書式が不十分」
ITEM_WORDS = ["食料", "住居", "光熱", "水道", "家具", "家事", "被服", "履物", "医療", "交通", "通信", "教育", "教養", "娯楽", "その他"]
CV_WORDS = r"変動係数|ばらつき|バラツキ|散らば|散布|差が大き|差が小さ"
SRC_DATA = [[271615, 299308, 356133, 331570, 265898], [72099, 87696, 92838, 88533, 80720], [31426, 17693, 17676, 17335, 17150],
            [18474, 21107, 23663, 23933, 23678], [13506, 12863, 12662, 13564, 12027], [10346, 12608, 13545, 10283, 6569],
            [12706, 12277, 13364, 15002, 17460], [43814, 43874, 53223, 49706, 32137], [3853, 15175, 34156, 13621, 599],
            [27364, 34197, 32988, 31562, 24963], [38026, 41818, 62019, 68031, 50595]]   # 配布した元データ(家計調査2024年、11費目×5年齢階級)
NOTE_PREFIX = "計算結果を読み取るときは"    # 模範解答の「注意事項」のテキストボックス(考察に数えない)

def register(api):
    LABELS, COMMENTS, nl = api.LABELS, api.COMMENTS, api.norm_label
    LABELS.update({
        "C2_NO_TABLE": "平均・中央値・標準偏差・変動係数の表が見つからない/未完成",
        "C2_HARDCODED": "関数・数式を使わず数値を入力している",
        "C2_WRONG_VALUE": "計算結果が誤り(関数の範囲の指定など)",
        "C2_WRONG_CV": "変動係数の式が誤り",
        "C2_WRONG_FUNC": "使っている関数が違う",
        "C2_STDEV_P": "標準偏差に STDEV.P(母集団)を使っている",
        "C2_NOT_FUNC": "関数を使わずに計算している(SUM/5 など。値は正しい)",
        "C2_INCOMPLETE": "表に空欄がある",
        "C2_FMT_COMMA": "桁区切りスタイルが設定されていない",
        "C2_FMT_DEC": "小数点以下の表示桁数が違う",
        "C2_NO_ESSAY": "考察がない",
        "C2_SHORT_ESSAY": "考察が短い",
        "C2_ESSAY_NO_CV": "考察で変動係数(ばらつき)に触れていない",
        "C2_ESSAY_FEW_ITEMS": "考察で具体的な費目に触れていない",
        "C2_ESSAY_IN_CELL": "考察をテキストボックスでなくセルに書いている(注意のみ)",
        "C2_MULTI_SHEET": "シートが複数ある",
        "C2_NO_RANGE": "範囲が計算されていない",
        "C2_MISSING_COL": "計算していない項目がある",
        "C2_OLD_DATA": "配布したデータと違うデータで解答している(以前の課題のファイルなど)",
        "C2_SRC_MODIFIED": "元データが書き換えられている(注意のみ)",
    })
    COMMENTS.update({
        "C2_NO_TABLE": "平均・中央値・標準偏差・変動係数の表が見つかりませんでした。",
        "C2_HARDCODED": "数値が手入力されています。AVERAGE、MEDIAN、STDEV.S などの関数を使って計算しましょう。",
        "C2_WRONG_VALUE": "{items}の計算結果が正しくない箇所があります。関数の範囲(例: B5:F5)を確認しましょう。",
        "C2_WRONG_FUNC": "{items}の関数が違うため、正しく計算されていません。平均はAVERAGE、中央値はMEDIAN、範囲はMAX−MIN、標準偏差はSTDEV.Sで求めます。",
        "C2_WRONG_CV": "変動係数の計算結果が正しくない箇所があります。変動係数は「標準偏差÷平均」で求めます。",
        "C2_STDEV_P": "標準偏差は STDEV.P ではなく STDEV.S を使いましょう(このデータは標本として扱います)。",
        "C2_NOT_FUNC": "SUM(…)/5 のような計算ではなく、AVERAGE などの関数を使いましょう。",
        "C2_INCOMPLETE": "表に空欄があります。すべての費目について計算しましょう。",
        "C2_FMT_COMMA": "{items}に桁区切りスタイルが設定されていません。",
        "C2_FMT_DEC": "{items}の小数点以下の表示桁数を確認しましょう(平均・中央値・標準偏差は小数第1位、変動係数は小数第1位または第2位まで)。",
        "C2_NO_ESSAY": "考察が書かれていません。計算結果から読み取れることを、テキストボックスにまとめましょう。",
        "C2_SHORT_ESSAY": "考察をもう少し詳しく書いてみましょう。",
        "C2_ESSAY_NO_CV": "変動係数から、年齢階級によるばらつきが大きい費目・小さい費目を読み取ってみましょう。",
        "C2_ESSAY_FEW_ITEMS": "「教育」「住居」など具体的な費目を挙げて、元のデータの傾向(年齢とともに増える・減る、山形など)にも触れましょう。",
        "C2_ESSAY_IN_CELL": "考察はセルではなく、テキストボックスに書きましょう(今回はセルの文章を考察として評価しています)。",
        "C2_MULTI_SHEET": "提出するシートは、課題のシート1枚だけにしましょう。",
        "C2_NO_RANGE": "範囲(最大値−最小値)も計算しましょう。",
        "C2_MISSING_COL": "{items}が見つかりませんでした。平均・中央値・範囲・標準偏差・変動係数をすべて計算し、見出しを付けましょう。",
        "C2_OLD_DATA": "解答に使われているデータが、今回配布した課題のデータ(2024年、5つの年齢階級)と違います。今回の課題ファイルを使って、もう一度解答してください。",
        "C2_SRC_MODIFIED": "元のデータが書き換えられているところがあります。元データは変更しないでください。",
    })

    def num(v):
        return isinstance(v, (int, float)) and not isinstance(v, bool)

    def stat_key(v):
        """見出しの文字から、どの統計量の列かを判定する(「標準偏差値」「平均値」のような表記も可)。"""
        t = nl(v)
        if not t or len(t) > 12 or "四分位" in t: return None      # 四分位範囲・四分位数は課題で求めていないので見ない(「範囲」と取り違えない)
        for name, k in STATS.items():
            if name in t: return k
        return None

    def find_layout(ws, wsv):
        """「平均」「中央値」「標準偏差」「変動係数」の見出しの行と列、データ(5つの年齢階級)の列、費目の行を探す。
        表の位置がずれていても(表側がB列から始まる、など)、見出しのすぐ左の5列をデータとみなす。データや表側が数式(他のブックの参照など)でもよい。"""
        for r in range(1, 41):
            cols = {}
            for c in range(1, 40):
                v = ws.cell(r, c).value
                k = stat_key(v) if isinstance(v, str) and not v.startswith("=") else None
                if k and k not in cols: cols[k] = c
            if len(cols) >= 3:
                c0 = min(cols.values()); dcols = list(range(c0 - 5, c0))
                if dcols[0] < 2: continue
                rows = [rr for rr in range(r + 1, r + 40) if nl(wsv.cell(rr, dcols[0] - 1).value) and
                        all(num(wsv.cell(rr, c).value) for c in dcols)]
                return r, cols, rows, dcols
        return None, {}, [], []

    def fmt_parts(nf, kind):
        """(桁区切りがOKか, 小数の桁数がOKか)。変動係数は桁区切りを問わず、小数第1位か第2位ならOK。"""
        t = re.sub(r'_.|"[^"]*"|\\.|\*.|\[[^\]]*\]|\s', "", str(nf or "General").split(";")[0])
        dec = len(t.split(".")[1]) if re.fullmatch(r"[#,0]*0\.0+", t) else (0 if re.fullmatch(r"[#,0]*0", t) else None)
        if kind == "cv": return True, dec in (1, 2)
        return "#,##0" in t, dec == 1

    def classify(cell, val, kind, exp, exp_p):
        v = cell.value
        v = getattr(v, "text", v)                # 配列数式(ArrayFormula)は、式の文字列を使う
        if num(v): return "hard"
        if not (isinstance(v, str) and v.startswith("=")): return "empty"
        f = re.sub(r"\s+", "", v).upper()
        if kind == "sd" and re.search(r"STDEV\.?P\b|STDEVP\(", f): return "stdev_p"
        if not num(val):                         # 計算結果が保存されていない(Excel以外で保存など) → 式の形だけで判定
            return "ok" if kind not in FUNC or re.search(FUNC[kind], f) else "not_func"
        tol = 0.0006 if kind == "cv" else max(0.06, abs(exp) * 1e-6)
        ok = exp is not None and (abs(val - exp) <= tol or kind == "cv" and abs(val - exp * 100) <= 0.06)   # 変動係数は百分率(×100)でもよい
        if kind == "cv" and not ok and exp_p is not None and abs(val - exp_p) <= 0.0006: return "stdev_p"   # 標準偏差の列がSTDEV.P
        if not ok: return "wrong_func" if kind in FUNC and not re.search(FUNC[kind], f) else "wrong"   # 関数そのものが違う(AVEDEVなど)か、範囲などの誤りか
        if kind in FUNC and not re.search(FUNC[kind], f): return "not_func"
        return "ok"                              # 変動係数は =I5/G5 でも =STDEV.S(…)/AVERAGE(…) でもよい

    def sheet_textboxes(path, title):
        """指定したシートに置かれたテキストボックスの文字だけを読む(ブック全体を提出した答案で、ほかのシートの文章を拾わないため)。"""
        import posixpath
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
            rd = lambda n: z.read(n).decode("utf8", "ignore") if n in names else ""
            rels = lambda base: {m.group(1): posixpath.normpath(posixpath.join(posixpath.dirname(base), m.group(2)))
                                 for m in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rd(posixpath.join(posixpath.dirname(base), "_rels", posixpath.basename(base) + ".rels")))}
            wbx = rd("xl/workbook.xml")
            rid = next((m.group(2) for m in re.finditer(r'<sheet [^>]*name="([^"]*)"[^>]*r:id="([^"]+)"', wbx) if m.group(1) == title), None)
            sheet = rels("xl/workbook.xml").get(rid) if rid else None
            if not sheet: return api.textboxes(path)
            out = []
            for d in [t for t in rels(sheet).values() if "drawings/drawing" in t]:
                for sp in re.split(r"<xdr:(?:twoCellAnchor|oneCellAnchor|absoluteAnchor)", rd(d))[1:]:
                    t = "\n".join("".join(re.findall(r"<a:t>(.*?)</a:t>", q, flags=re.S)) for q in re.findall(r"<a:p>.*?</a:p>", sp, flags=re.S)).strip()
                    if t: out.append(t)
            return "\n".join(out)

    def essay_of(path, ws, head_row):
        tb = "\n".join(p for p in sheet_textboxes(path, ws.title).split("\n") if not nl(p).startswith(NOTE_PREFIX))
        cells = []
        skip_rows = set(range(1, head_row + 1))
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str) or v.startswith("="): continue
                t = nl(v)
                if len(t) < 2 or c.row in skip_rows and "。" not in t: continue
                if c.column == 1 and "。" not in t: continue                        # 表側の費目名など
                if re.match(r"(単位|出所|出典|資料|注[:：\s1-9(（]|注$)", t): continue                 # 単位・出所・注の記載
                if re.fullmatch(r"(考察|考察の例|注意事項)[:：]?", t): continue
                if t in STATS: continue
                cells.append(v.strip())
        cell_text = "\n".join(cells)
        if len(nl(cell_text)) >= MIN_ESSAY_NONE and len(nl(tb)) < MIN_ESSAY:
            return (tb + "\n" + cell_text).strip(), True
        return tb, False

    @api.checker(2)
    def check2(path, ref):
        wb = load_workbook(path); wbv = load_workbook(path, data_only=True)
        iss = []
        if len(wb.worksheets) > 1: iss.append("C2_MULTI_SHEET")
        ws, wsv = wb.worksheets[0], wbv.worksheets[0]
        h, cols, rows, dcols = find_layout(ws, wsv)
        if h is None or not rows:                      # 1枚目にないときは、ほかのシートから表を探す(ブック全体を提出した答案など)
            for a, b in zip(wb.worksheets[1:], wbv.worksheets[1:]):
                lay = find_layout(a, b)
                if lay[0] is not None and lay[2]: ws, wsv = a, b; h, cols, rows, dcols = lay; break
        text, in_cell = essay_of(path, ws, h or 0)
        if in_cell: iss.append("C2_ESSAY_IN_CELL")
        n = len(nl(text))
        if n < MIN_ESSAY_NONE: essay = "なし"; iss.append("C2_NO_ESSAY")
        else:
            lack = []
            if n < MIN_ESSAY: lack.append("C2_SHORT_ESSAY")
            if not re.search(CV_WORDS, text): lack.append("C2_ESSAY_NO_CV")
            if len({w for w in ITEM_WORDS if w in text}) < NEED_ITEMS: lack.append("C2_ESSAY_FEW_ITEMS")
            essay = "不十分" if lack else "妥当"; iss += lack
        if h is None or len(cols) < 4 or not rows:
            return "評価外", iss + ["C2_NO_TABLE"], {"考察文字数": n}
        ages = [c for c in range(1, min(cols.values())) if isinstance(wsv.cell(h, c).value, str) and re.search(r"歳|～|~", wsv.cell(h, c).value)]
        diff = sum(1 for k, r in enumerate(rows[:len(SRC_DATA)]) for j, c in enumerate(dcols) if wsv.cell(r, c).value != SRC_DATA[k][j])
        if len(ages) > 5 or len(rows) < len(SRC_DATA) - 2 or diff > len(SRC_DATA) * 5 / 2:     # 年齢階級の数が違う、データの大半が違う → 以前の課題など
            return "評価外", [c for c in iss if c == "C2_MULTI_SHEET"] + ["C2_OLD_DATA"], {"考察文字数": n}
        if diff: iss.append("C2_SRC_MODIFIED")
        # 見出しの誤字(「変動偏差」など)で見つからなかった列は、表の右側の列の値から推定する
        def expv(r, kind):
            d = [wsv.cell(r, c).value for c in dcols]; m = statistics.mean(d); sd = statistics.stdev(d)
            return {"avg": m, "med": statistics.median(d), "rng": max(d) - min(d), "sd": sd, "cv": sd / m if m else None}[kind]
        c0 = min(cols.values())
        for kind in [k for k in STATS.values() if k not in cols]:
            for c in range(c0, c0 + 10):
                if c in cols.values(): continue
                vals = [wsv.cell(r, c).value for r in rows]
                hit = sum(1 for r, v in zip(rows, vals) if num(v) and (e := expv(r, kind)) is not None and
                          (abs(v - e) <= max(0.06, abs(e) * 1e-6) if kind != "cv" else (abs(v - e) <= 0.0006 or abs(v - e * 100) <= 0.06)))
                if hit >= len(rows) * 0.8: cols[kind] = c; break
        cols = dict(sorted(cols.items(), key=lambda kv: kv[1]))
        cnt, filled, comma_ng, dec_ng, wrong_at = Counter(), Counter(), Counter(), Counter(), Counter()
        NAME = {v: k for k, v in STATS.items()}
        for r in rows:
            data = [wsv.cell(r, c).value for c in dcols]
            m = statistics.mean(data); sd = statistics.stdev(data); sdp = statistics.pstdev(data)
            exp = {"avg": m, "med": statistics.median(data), "rng": max(data) - min(data), "sd": sd, "cv": sd / m if m else None}
            for kind, c in cols.items():
                cell = ws.cell(r, c)
                k = classify(cell, wsv.cell(r, c).value, kind, exp[kind], sdp / m if (kind == "cv" and m) else None)
                if kind == "sd" and k == "wrong" and abs(wsv.cell(r, c).value - sdp) <= 0.06: k = "stdev_p"
                if kind == "cv" and k in ("wrong", "stdev_p") and "sd" in cols and "avg" in cols:    # 学生自身の標準偏差÷平均と合っていれば、式は正しい(平均などの誤りを二重に数えない)
                    s_, a_, v_ = wsv.cell(r, cols["sd"]).value, wsv.cell(r, cols["avg"]).value, wsv.cell(r, c).value
                    if num(s_) and num(a_) and a_ and (abs(v_ - s_ / a_) <= 0.0006 or abs(v_ - s_ / a_ * 100) <= 0.06): k = "ok"
                if k == "wrong_func": wrong_at["func_" + kind] += 1; k = "wrong"
                elif k == "wrong": wrong_at[kind] += 1
                cnt[k] += 1
                if k != "empty":
                    a, d = fmt_parts(cell.number_format, kind)
                    filled[kind] += 1; comma_ng[kind] += not a; dec_ng[kind] += not d
        total = len(rows) * len(cols)
        if (total - cnt["empty"]) / total < 0.2:
            return "評価外", iss + ["C2_NO_TABLE"], dict(cnt) | {"考察文字数": n}
        if cnt["empty"]: iss.append("C2_INCOMPLETE")
        if cnt["hard"] > HAND_TYPED_MAX:
            calc = "式なし"; iss.append("C2_HARDCODED")
        elif (cnt["wrong"] + cnt["empty"]) / total > 0.05:
            calc = "値誤り"
            wf = [NAME[k] for k in cols if wrong_at["func_" + k]]
            w = [NAME[k] for k in cols if wrong_at[k] and k != "cv" and not wrong_at["func_" + k]]
            if wf: iss.append("C2_WRONG_FUNC@" + ",".join(wf))
            if w or not (wrong_at["cv"] or wf): iss.append("C2_WRONG_VALUE" + ("@" + ",".join(w) if w else ""))
            if wrong_at["cv"]: iss.append("C2_WRONG_CV")
        else:
            calc = "正"
        if cnt["stdev_p"]: iss.append("C2_STDEV_P")
        if cnt["not_func"]: iss.append("C2_NOT_FUNC")
        if "rng" not in cols: iss.append("C2_NO_RANGE")
        miss = [NAME[k] for k in ("avg", "med", "sd", "cv") if k not in cols]
        if miss and calc == "正": calc = "値誤り"
        if miss: iss.append("C2_MISSING_COL@" + ",".join(miss))
        bad = lambda ng: [NAME[k] for k in cols if filled[k] and ng[k] / filled[k] > 1 - FMT_OK_RATIO]   # 項目ごとに、9割以上のセルで正しければOK
        bc, bd = bad(comma_ng), bad(dec_ng)
        if bc: iss.append("C2_FMT_COMMA@" + ",".join(bc))
        if bd: iss.append("C2_FMT_DEC@" + ",".join(bd))
        fmt = not (bc or bd)
        ESSAY = ("C2_NO_ESSAY", "C2_SHORT_ESSAY", "C2_ESSAY_NO_CV", "C2_ESSAY_FEW_ITEMS", "C2_ESSAY_IN_CELL")
        iss = [c for c in iss if not c.startswith(ESSAY)] + [c for c in iss if c.startswith(ESSAY)]     # コメントは計算・書式が先、考察が後
        # S〜D: 手入力→D。計算結果の誤り(関数の範囲の誤りなど)→B(考察なしはC)。考察なし→B。
        # それ以外はSとし、軽い不足(STDEV.P、関数を使わない計算、書式、範囲なし、考察が明らかに短い)が1つでもあればA。
        # 考察の語句の不足(変動係数・費目に触れていない等)は、AI判定ではないのでコメントのみで減点しない
        minor = cnt["stdev_p"] or cnt["not_func"] or not fmt or "rng" not in cols or (essay != "なし" and n < ESSAY_DEDUCT)
        if calc == "式なし": grade = "D"
        elif calc == "値誤り": grade = "C" if essay == "なし" else "B"
        elif essay == "なし": grade = "B"
        else: grade = "A" if minor else "S"
        return grade, iss, {f"n_{k}": v for k, v in cnt.items()} | {"考察文字数": n}
