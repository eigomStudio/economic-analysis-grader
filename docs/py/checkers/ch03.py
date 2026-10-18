"""第3章 章末課題(輸出額・輸入額の寄与度・寄与率と、寄与度のグラフ)のチェッカー。
grader.py と同じフォルダの checkers/ に置くと読み込まれる(Webアプリでも publish.py でコピーされる)。

3つの項目を「できている/不十分/できていない」で判定する。
- 計算: 輸出・輸入の寄与度・寄与率(4つの表)の計算結果が正しいか(絶対参照は問わない。使っていなければ注意のみ)
- グラフ: 輸出・輸入の寄与度のグラフ(内訳は積み上げ縦棒、総額は折れ線、横軸のラベルは「下端/左端」、凡例は右)
- 考察: 字数と語句(輸出・輸入の両方、寄与度・寄与率への言及、具体的な品目名)
評価: できていない項目が 0→S、1→A、2→C、3→D。「不十分」が1つでもあれば1段階下げる(下限D)。
"""
import re
import zipfile
from collections import Counter
from openpyxl import load_workbook

ITEMS = ["食料品及び動物", "飲料及びたばこ", "食料に適さない原材料", "鉱物性燃料", "動植物性油脂",
         "化学製品", "原料別製品", "機械類及び輸送用機器", "雑製品", "特殊取扱品"]
TABLES = {"exp_deg": "輸出額の寄与度", "imp_deg": "輸入額の寄与度", "exp_rat": "輸出額の寄与率", "imp_rat": "輸入額の寄与率"}
CALC_OK, CALC_HALF = 0.95, 0.5          # 正しいセルの割合: 95%以上=できている、50%以上=不十分、未満=できていない
MIN_ESSAY, MIN_ESSAY_NONE, NEED_ITEMS = 150, 10, 2
ITEM_WORDS = ["食料品", "飲料", "原材料", "鉱物性燃料", "燃料", "油脂", "化学製品", "原料別製品", "機械", "輸送用機器", "雑製品", "特殊取扱品"]
NOTE_PREFIX = "寄与度・寄与率の分析は"     # 模範解答の「注意事項」のテキストボックス(考察に数えない)
ORDER = ["S", "A", "B", "C", "D"]

def register(api):
    LABELS, COMMENTS, nl = api.LABELS, api.COMMENTS, api.norm_label
    LABELS.update({
        "C3_HARDCODED": "寄与度・寄与率を数式でなく数値で入力している",
        "C3_NO_TABLE": "寄与度・寄与率の表が見つからない",
        "C3_WRONG": "寄与度・寄与率の計算結果に誤りがある",
        "C3_MISSING": "寄与度・寄与率の表が不足している",
        "C3_NO_ABS": "絶対参照を使っていない(注意のみ)",
        "C3_NO_CHART": "寄与度のグラフがない",
        "C3_ONE_CHART": "寄与度のグラフが1つしかない",
        "C3_CHART_DATA": "グラフのデータ範囲が寄与度の表になっていない",
        "C3_CHART_BAR": "内訳が積み上げ縦棒になっていない",
        "C3_CHART_LINE": "総額が折れ線になっていない",
        "C3_CHART_AXIS": "横軸のラベルの位置が「下端/左端」になっていない",
        "C3_CHART_LEGEND": "凡例の位置が右になっていない",
        "C3_NO_ESSAY": "考察がない",
        "C3_SHORT_ESSAY": "考察が短い",
        "C3_ESSAY_BOTH": "考察で輸出・輸入の両方に触れていない",
        "C3_ESSAY_KIYO": "考察で寄与度・寄与率に触れていない",
        "C3_ESSAY_ITEMS": "考察で具体的な品目に触れていない",
        "C3_ESSAY_IN_CELL": "考察をテキストボックスでなくセルに書いている(注意のみ)",
    })
    COMMENTS.update({
        "C3_HARDCODED": "寄与度・寄与率が数値で入力されています。セル参照を使った数式で計算しましょう。",
        "C3_NO_TABLE": "寄与度・寄与率の表が見つかりませんでした。輸出・輸入それぞれの寄与度と寄与率の表を作成しましょう。",
        "C3_WRONG": "{items}の計算結果が正しくない箇所があります。寄与度は「(今年の値−前年の値)÷前年の総額×100」、寄与率は「(今年の値−前年の値)÷(今年の総額−前年の総額)×100」です。",
        "C3_MISSING": "{items}の表が見つかりませんでした。輸出・輸入それぞれについて、寄与度と寄与率の両方を計算しましょう。",
        "C3_NO_ABS": "総額の行を B$13 のように行を固定した絶対参照にすると、1つの式をコピーするだけで表全体を計算できます。",
        "C3_NO_CHART": "寄与度のグラフがありません。輸出・輸入それぞれの寄与度を、グラフにしましょう。",
        "C3_ONE_CHART": "寄与度のグラフが1つしかありません。輸出と輸入の両方のグラフを作成しましょう。",
        "C3_CHART_DATA": "グラフには、寄与度の表のデータを使いましょう。",
        "C3_CHART_BAR": "{items}のグラフで、品目の寄与度を積み上げ縦棒グラフにしましょう。",
        "C3_CHART_LINE": "{items}のグラフで、総額(全体の伸び率)を折れ線グラフにしましょう。",
        "C3_CHART_AXIS": "{items}のグラフで、横軸のラベルの位置を「下端/左端」にすると、マイナスの棒と重ならず読みやすくなります。",
        "C3_CHART_LEGEND": "{items}のグラフで、凡例の位置を右にしましょう。",
        "C3_NO_ESSAY": "考察が書かれていません。寄与度・寄与率の計算結果から読み取れることを、テキストボックスにまとめましょう。",
        "C3_SHORT_ESSAY": "考察をもう少し詳しく書いてみましょう。",
        "C3_ESSAY_BOTH": "輸出と輸入の両方の動きについて考察しましょう。",
        "C3_ESSAY_KIYO": "どの品目が総額の増減に大きく寄与したか(寄与度・寄与率)を読み取ってみましょう。",
        "C3_ESSAY_ITEMS": "「機械類及び輸送用機器」「鉱物性燃料」など、具体的な品目を挙げて考察しましょう。",
        "C3_ESSAY_IN_CELL": "考察はセルではなく、テキストボックスに書きましょう(今回はセルの文章を考察として評価しています)。",
    })
    ITEMN = {nl(x) for x in ITEMS}

    def num(v): return isinstance(v, (int, float)) and not isinstance(v, bool)

    def is_item(v):
        t = nl(v)
        return t in ITEMN or bool(re.fullmatch(r"(輸出|輸入)?総額|合計|総計", t))

    def blocks(ws, r0, r1):
        """A列に品目名が10行+総額が並ぶまとまり(表)を探す。[(先頭行, 行数)]"""
        out, r = [], r0
        while r <= r1:
            if nl(ws.cell(r, 1).value) == nl(ITEMS[0]):
                k = 0
                while k < 11 and is_item(ws.cell(r + k, 1).value): k += 1
                if k >= 10: out.append((r, k)); r += k; continue
            r += 1
        return out

    def year_cols(wsv, ws, hr):
        """表頭(年)の行から、年→列の対応を作る。"""
        m = {}
        for c in range(2, 40):
            v = wsv.cell(hr, c).value
            if num(v) and 1990 <= v <= 2100: m[int(v)] = c
        return m

    def essay_of(path, ws):
        boxes = []                                     # テキストボックスごとに読み、「注意事項」のボックスは丸ごと除く
        with zipfile.ZipFile(path) as z:
            for nm in z.namelist():
                if nm.startswith("xl/drawings/drawing") and nm.endswith(".xml"):
                    for sp in re.split(r"<xdr:(?:twoCellAnchor|oneCellAnchor|absoluteAnchor)", z.read(nm).decode("utf8", "ignore"))[1:]:
                        t = "\n".join("".join(re.findall(r"<a:t>(.*?)</a:t>", q, flags=re.S)) for q in re.findall(r"<a:p>.*?</a:p>", sp, flags=re.S)).strip()
                        if t and not nl(t).startswith(nl(NOTE_PREFIX)): boxes.append(t)
        tb = "\n".join(boxes)
        cells = []
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str) or v.startswith("="): continue
                t = nl(v)
                if len(t) < 2 or "。" not in t and (c.column == 1 or len(t) < 20): continue      # 表題・品目名・見出しなど
                if re.match(r"(単位|出所|出典|資料|注[:：\s1-9(（]|注$)", t): continue
                cells.append(v.strip())
        cell_text = "\n".join(cells)
        if len(nl(cell_text)) >= MIN_ESSAY_NONE and len(nl(tb)) < MIN_ESSAY:
            return (tb + "\n" + cell_text).strip(), True
        return tb, False

    def charts(path):
        """ブック内のグラフの、棒/折れ線の系列(値の範囲)と、横軸ラベル・凡例の位置。"""
        out = []
        with zipfile.ZipFile(path) as z:
            for n in sorted(x for x in z.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", x)):
                x = z.read(n).decode("utf8", "ignore")
                ser = lambda body: re.findall(r"<c:val>.*?<c:f>([^<]+)</c:f>", body, re.S)
                bar = [(g, ser(b)) for b, g in [(m.group(1), m.group(2)) for m in re.finditer(r"<c:barChart>(.*?<c:grouping val=\"(\w+)\".*?)</c:barChart>", x, re.S)]]
                bdir = re.findall(r"<c:barDir val=\"(\w+)\"", x)
                line = [ser(m.group(1)) for m in re.finditer(r"<c:lineChart>(.*?)</c:lineChart>", x, re.S)]
                cat = re.search(r"<c:catAx>(.*?)</c:catAx>", x, re.S)
                tick = re.search(r"<c:tickLblPos val=\"(\w+)\"", cat.group(1)).group(1) if cat and "<c:tickLblPos" in cat.group(1) else "nextTo"
                lg = re.search(r"<c:legend>(.*?)</c:legend>", x, re.S)
                pos = (re.search(r"<c:legendPos val=\"(\w+)\"", lg.group(1)) or [None, "r"])[1] if lg else None
                out.append({"bar": bar, "bardir": bdir, "line": [s for l in line for s in l], "tick": tick, "legend": pos})
        return out

    def rows_of(ref):
        m = re.search(r"\$?[A-Z]{1,3}\$?(\d+)(?::\$?[A-Z]{1,3}\$?(\d+))?$", ref)
        return {int(m.group(1)), int(m.group(2) or m.group(1))} if m else set()

    @api.checker(3)
    def check3(path, ref):
        wb = load_workbook(path); ws = wb.worksheets[0]
        wsv = load_workbook(path, data_only=True).worksheets[0]
        iss, det = [], {}
        # ── 元データ(輸出・輸入の表)
        src = blocks(ws, 1, 40)
        if len(src) < 2:
            return "評価外", ["C3_NO_TABLE"], {}
        (er, _), (ir, _) = src[0], src[1]
        sy = {"exp": year_cols(wsv, ws, er - 1), "imp": year_cols(wsv, ws, ir - 1)}
        srow = {"exp": er, "imp": ir}
        def sval(kind, k, y):
            c = sy[kind].get(y); v = wsv.cell(srow[kind] + k, c).value if c else None
            return v if num(v) else None
        def expect(t, k, y):
            kind = t[:3]; a, b = sval(kind, k, y), sval(kind, k, y - 1); ta, tb = sval(kind, 10, y), sval(kind, 10, y - 1)
            if None in (a, b, ta, tb): return None
            if t.endswith("deg"): return (a - b) / tb * 100 if tb else None
            return (a - b) / (ta - tb) * 100 if ta != tb else None
        # ── 学生の寄与度・寄与率の表
        found, cnt = {}, Counter()
        for r0, n in blocks(ws, srow["imp"] + 11, ws.max_row):
            ycol = year_cols(wsv, ws, r0 - 1) or {2010 + i: 2 + i for i in range(16)}
            best = None
            for t in TABLES:
                if t in found: continue
                ok = sum(1 for k in range(n) for y, c in ycol.items()
                         if (e := expect(t, k, y)) is not None and num(v := wsv.cell(r0 + k, c).value) and (abs(v - e) <= 0.006 or abs(v - e / 100) <= 0.00006))
                if best is None or ok > best[1]: best = (t, ok)
            if best is None: continue
            t = best[0]
            if best[1] == 0:                                   # 値が1つも合わない表は、見出しから種類を推測する
                title = "".join(nl(ws.cell(r, 1).value) for r in range(r0 - 3, r0))
                guess = [x for x in TABLES if x not in found and TABLES[x][:2] in title and TABLES[x][-3:] in title]
                if guess: t = guess[0]
            found[t] = (r0, n, ycol)
            for k in range(n):
                for y, c in ycol.items():
                    e, cell, v = expect(t, k, y), ws.cell(r0 + k, c), wsv.cell(r0 + k, c).value
                    if e is None: continue
                    if num(cell.value): cnt["hard"] += 1; continue
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")): cnt["empty"] += 1; continue
                    good = num(v) and (abs(v - e) <= 0.006 or abs(v - e / 100) <= 0.00006)
                    cnt["ok" if good else "wrong"] += 1; cnt[f"wrong_{t}"] += not good
                    cnt["abs"] += "$" in cell.value
        total = len(TABLES) * 11 * 16
        missing = [TABLES[t] for t in TABLES if t not in found]
        if cnt["hard"]: calc = "NG"; iss.append("C3_HARDCODED")
        elif not found: calc = "NG"; iss.append("C3_NO_TABLE")
        else:
            ratio = cnt["ok"] / total
            calc = "OK" if ratio >= CALC_OK else "HALF" if ratio >= CALC_HALF else "NG"
            w = [TABLES[t] for t in TABLES if cnt[f"wrong_{t}"]]
            if w: iss.append("C3_WRONG@" + ",".join(w))
            if missing: iss.append("C3_MISSING@" + ",".join(missing))
            if cnt["ok"] + cnt["wrong"] and cnt["abs"] / (cnt["ok"] + cnt["wrong"]) < 0.5: iss.append("C3_NO_ABS")
        # ── グラフ
        degrows = {t: set(range(found[t][0], found[t][0] + found[t][1])) for t in ("exp_deg", "imp_deg") if t in found}
        ch_ok = {}
        probs = Counter()
        for ch in charts(path):
            refs = [s for _, ss in ch["bar"] for s in ss] + ch["line"]
            rows = set().union(*(rows_of(s) for s in refs)) if refs else set()
            kind = next((t for t, rr in degrows.items() if rows and rows <= rr), None)
            if kind is None:
                kind = next((t for t in ("exp_deg", "imp_deg") if t not in ch_ok), "exp_deg"); probs["data"] += 1
            name = "輸出" if kind == "exp_deg" else "輸入"
            p = set()
            stacked = [ss for g, ss in ch["bar"] if g == "stacked"]
            if not stacked or "bar" in ch["bardir"] or sum(len(s) for s in stacked) < 8: p.add("bar")
            line_rows = set().union(*(rows_of(s) for s in ch["line"])) if ch["line"] else set()
            tot = (found[kind][0] + found[kind][1] - 1) if kind in found else None
            if not ch["line"] or (tot and tot not in line_rows): p.add("line")
            if ch["tick"] != "low": p.add("axis")
            if ch["legend"] != "r": p.add("legend")
            ch_ok[kind] = p
            for x in p: probs[(x, name)] += 1
        if not ch_ok: chart = "NG"; iss.append("C3_NO_CHART")
        else:
            chart = "OK"
            if len(ch_ok) < 2: chart = "HALF"; iss.append("C3_ONE_CHART")
            if probs["data"]: chart = "HALF"; iss.append("C3_CHART_DATA")
            for x, code in (("bar", "C3_CHART_BAR"), ("line", "C3_CHART_LINE"), ("axis", "C3_CHART_AXIS"), ("legend", "C3_CHART_LEGEND")):
                names = [nm for nm in ("輸出", "輸入") if probs[(x, nm)]]
                if names: chart = "HALF"; iss.append(code + "@" + ",".join(f"{nm}の寄与度" for nm in names))
        # ── 考察
        text, in_cell = essay_of(path, ws)
        n = len(nl(text))
        ess_codes = ["C3_ESSAY_IN_CELL"] if in_cell else []
        if n < MIN_ESSAY_NONE: essay = "NG"; ess_codes.append("C3_NO_ESSAY")
        else:
            lack = []
            if n < MIN_ESSAY: lack.append("C3_SHORT_ESSAY")
            if not ("輸出" in text and "輸入" in text): lack.append("C3_ESSAY_BOTH")
            if "寄与" not in text: lack.append("C3_ESSAY_KIYO")
            if len({w for w in ITEM_WORDS if w in text}) < NEED_ITEMS: lack.append("C3_ESSAY_ITEMS")
            essay = "HALF" if lack else "OK"; ess_codes += lack
        iss += ess_codes                                       # コメントは計算 → グラフ → 考察の順
        # ── 評価
        st = [calc, chart, essay]
        ng = st.count("NG")
        i = {0: 0, 1: 1, 2: 3, 3: 4}[ng] + (1 if "HALF" in st else 0)
        grade = ORDER[min(i, 4)]
        J = {"OK": "できている", "HALF": "不十分", "NG": "できていない"}
        det = {"計算": J[calc], "グラフ": J[chart], "考察": J[essay], "考察文字数": n} | {f"n_{k}": v for k, v in cnt.items() if not k.startswith("wrong_")}
        return grade, iss, det
