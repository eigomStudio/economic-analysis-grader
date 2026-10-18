"""第4章 章末課題(OECD諸国の女性労働力率と合計特殊出生率: 散布図と相関係数)のチェッカー。
grader.py と同じフォルダの checkers/ に置くと読み込まれる(Webアプリでも publish.py でコピーされる)。

3つの項目を「できている/不十分/できていない」で判定する(第3章と同じ基準)。
- 散布図: 1970・1985・2000年の3つ(横軸=女性労働力率、縦軸=合計特殊出生率)。縦・横の軸ラベル、横軸・縦軸の目盛りの最大値が3つでそろっているか
- 相関係数: 3つの年について、関数(CORREL または PEARSON)で正しく計算しているか
- 考察: 字数と語句。「労働力率を上げれば出生率が上がる」のような因果関係の読み方をしていれば「不十分」(要確認)
評価: できていない項目が 0→S、1→A、2→C、3→D。「不十分」が1つでもあれば1段階下げる(下限D)。
"""
import math
import re
import statistics
import zipfile
from openpyxl import load_workbook

YEARS = [1970, 1985, 2000]
MIN_ESSAY, MIN_ESSAY_NONE = 150, 10
NOTE_PREFIXES = ["グラフの変化を比較するために", "全てのグラフのメモリの最大値を揃えると"]   # 模範解答の注意点(考察に数えない)
ORDER = ["S", "A", "B", "C", "D"]
CAUSE = r"(労働力率|就業率|働く女性|女性の就業|女性が働)[^。]{0,30}(上げ|引き上げ|高め|増や|上が|高ま|増え|上昇)[^。]{0,6}(れば|ると|たら|ことで|ことによって|ため|ので)[^。]{0,30}出生率[^。]{0,10}(上が|高ま|上昇|増え|改善)"
NEGATE = r"わけではな|とは限らな|とはいえな|とは言えな|示していな|示すものではな|ではない|因果関係(は|が)(ない|な[いく]|示)"

def register(api):
    LABELS, COMMENTS, nl = api.LABELS, api.COMMENTS, api.norm_label
    LABELS.update({
        "C4_NO_CHART": "散布図がない",
        "C4_FEW_CHART": "散布図が3つ(1970・1985・2000年)そろっていない",
        "C4_NOT_SCATTER": "散布図になっていない、または縦軸・横軸が逆",
        "C4_NO_AXIS_TITLE": "軸ラベル(縦軸・横軸の名前)がない",
        "C4_SCALE": "軸の目盛りの最大値がそろっていない",
        "C4_NO_CORR": "相関係数が計算されていない",
        "C4_CORR_HARD": "相関係数を関数でなく数値で入力している",
        "C4_CORR_WRONG": "相関係数の値が誤り(範囲の指定など)",
        "C4_CORR_FEW": "相関係数が3つの年すべてで計算されていない",
        "C4_NO_ESSAY": "考察がない",
        "C4_SHORT_ESSAY": "考察が短い",
        "C4_ESSAY_CORR": "考察で相関(散布図の傾向や相関係数)に触れていない",
        "C4_ESSAY_YEARS": "考察で年ごとの変化に触れていない",
        "C4_CAUSAL": "相関を因果関係として読んでいる(要確認)",
        "C4_ESSAY_IN_CELL": "考察をテキストボックスでなくセルに書いている(注意のみ)",
    })
    COMMENTS.update({
        "C4_NO_CHART": "散布図がありません。1970年・1985年・2000年のそれぞれについて、女性労働力率(横軸)と合計特殊出生率(縦軸)の散布図を作成しましょう。",
        "C4_FEW_CHART": "散布図が3つそろっていません。1970年・1985年・2000年のすべてについて作成しましょう。",
        "C4_NOT_SCATTER": "{items}のグラフを確認しましょう。散布図で、横軸を女性労働力率、縦軸を合計特殊出生率にします。",
        "C4_NO_AXIS_TITLE": "{items}の散布図に、軸ラベル(縦軸・横軸の名前)を付けましょう。",
        "C4_SCALE": "3つの散布図で、{items}の目盛りの最大値をそろえましょう(横軸は90、縦軸は5.0)。そろえないと、年ごとの変化を正しく比較できません。",
        "C4_NO_CORR": "相関係数が見つかりませんでした。CORREL関数で、年ごとの相関係数を計算しましょう。",
        "C4_CORR_HARD": "相関係数が数値で入力されています。CORREL関数を使って計算しましょう。",
        "C4_CORR_WRONG": "{items}の相関係数が正しくありません。CORREL関数の2つの範囲(同じ年の女性労働力率と合計特殊出生率)を確認しましょう。",
        "C4_CORR_FEW": "{items}の相関係数が見つかりませんでした。3つの年すべてについて計算しましょう。",
        "C4_NO_ESSAY": "考察が書かれていません。散布図と相関係数から読み取れることを、テキストボックスにまとめましょう。",
        "C4_SHORT_ESSAY": "考察をもう少し詳しく書いてみましょう。",
        "C4_ESSAY_CORR": "散布図の傾き(右上がり・右下がり)や相関係数の値から、2つのデータの関係を読み取りましょう。",
        "C4_ESSAY_YEARS": "1970年・1985年・2000年で、関係がどう変わったかを比べてみましょう。",
        "C4_CAUSAL": "相関があることは、因果関係があることを示しているわけではありません。「女性の労働力率を上げれば出生率が上がる」とは言えない点に注意しましょう。正の相関の背景には、女性の労働力率が高い国々で、仕事と子育ての両立支援策が充実してきたことなどが考えられます。",
        "C4_ESSAY_IN_CELL": "考察はセルではなく、テキストボックスに書きましょう(今回はセルの文章を考察として評価しています)。",
    })

    def num(v): return isinstance(v, (int, float)) and not isinstance(v, bool)

    def data_rows(wsv):
        """B〜G列(女性労働力率3年分・出生率3年分)がすべて数値の行。"""
        return [r for r in range(1, 60) if all(num(wsv.cell(r, c).value) for c in range(2, 8)) and not all(1900 <= wsv.cell(r, c).value <= 2100 for c in range(2, 8))]

    def nice_max(m):
        """軸の最大値を自動にしたときにExcelが付ける最大値(最小値0のとき)。データの最大値に約5%の余白を足し、
        目盛り間隔(1・2・5×10^k、目盛りが10本以内)の倍数に切り上げる。例: 79.17→90、4.53→5.0"""
        if m <= 0: return 0
        m = m * 1.05
        step = 10 ** math.floor(math.log10(m / 10)); unit = next(u * step for u in (1, 2, 5, 10) if m / (u * step) <= 10)
        return round(math.ceil(m / unit - 1e-9) * unit, 10)

    def charts(path):
        out = []
        with zipfile.ZipFile(path) as z:
            for n in sorted(x for x in z.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", x)):
                x = z.read(n).decode("utf8", "ignore")
                xs = re.findall(r"<c:xVal>.*?<c:f>([^<]+)</c:f>", x, re.S)
                ys = re.findall(r"<c:yVal>.*?<c:f>([^<]+)</c:f>", x, re.S)
                axes = {}
                for ax in re.findall(r"<c:valAx>.*?</c:valAx>", x, re.S):
                    pos = (re.findall(r"<c:axPos val=\"(\w)\"", ax) or ["?"])[0]
                    mx = re.findall(r"<c:max val=\"([^\"]+)\"", ax)
                    axes["x" if pos in "bt" else "y"] = {"title": "<c:title>" in ax, "max": float(mx[0]) if mx else None}
                out.append({"scatter": "<c:scatterChart>" in x, "x": xs, "y": ys, "axes": axes})
        return out

    def col_of(ref):
        m = re.search(r"\$?([A-Z]{1,3})\$?\d+", ref.split("!")[-1])
        return m.group(1) if m else None

    def essay_of(path, ws):
        boxes = []
        with zipfile.ZipFile(path) as z:
            for nm in z.namelist():
                if nm.startswith("xl/drawings/drawing") and nm.endswith(".xml"):
                    for sp in re.split(r"<xdr:(?:twoCellAnchor|oneCellAnchor|absoluteAnchor)", z.read(nm).decode("utf8", "ignore"))[1:]:
                        t = "\n".join("".join(re.findall(r"<a:t>(.*?)</a:t>", q, flags=re.S)) for q in re.findall(r"<a:p>.*?</a:p>", sp, flags=re.S)).strip()
                        if t and not any(nl(t).startswith(nl(p)) for p in NOTE_PREFIXES): boxes.append(t)
        tb = "\n".join(boxes)
        cells = []
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str) or v.startswith("="): continue
                t = nl(v)
                if len(t) < 2 or "。" not in t and (c.column == 1 or len(t) < 20): continue
                if re.match(r"(単位|出所|出典|資料|注[:：\s1-9(（]|注$)", t): continue
                cells.append(v.strip())
        cell_text = "\n".join(cells)
        if len(nl(cell_text)) >= MIN_ESSAY_NONE and len(nl(tb)) < MIN_ESSAY:
            return (tb + "\n" + cell_text).strip(), True
        return tb, False

    @api.checker(4)
    def check4(path, ref):
        wb = load_workbook(path); ws = wb.worksheets[0]
        wsv = load_workbook(path, data_only=True).worksheets[0]
        rows = data_rows(wsv)
        if len(rows) < 5:
            return "評価外", ["C4_NO_CHART"], {}
        lab = {y: [wsv.cell(r, 2 + i).value for r in rows] for i, y in enumerate(YEARS)}
        tfr = {y: [wsv.cell(r, 5 + i).value for r in rows] for i, y in enumerate(YEARS)}
        exp = {y: statistics.correlation(lab[y], tfr[y]) for y in YEARS}
        iss = []
        # ── 散布図
        LCOL = {"B": 1970, "C": 1985, "D": 2000}; TCOL = {"E": 1970, "F": 1985, "G": 2000}
        found, bad, notitle, ymax, xmax = {}, [], [], {}, {}
        for ch in charts(path):
            xc = col_of(ch["x"][0]) if ch["x"] else None; yc = col_of(ch["y"][0]) if ch["y"] else None
            year = LCOL.get(xc) or TCOL.get(yc) or LCOL.get(yc) or TCOL.get(xc)
            if year is None or year in found: continue
            found[year] = ch
            if not ch["scatter"] or xc not in LCOL or yc not in TCOL: bad.append(f"{year}年")
            if not all(ch["axes"].get(k, {}).get("title") for k in ("x", "y")): notitle.append(f"{year}年")
            m = ch["axes"].get("y", {}).get("max")
            ymax[year] = m if m is not None else nice_max(max(tfr[year]))
            m = ch["axes"].get("x", {}).get("max")
            xmax[year] = m if m is not None else nice_max(max(lab[year]))     # 自動の場合は、Excelが付けるおよその最大値で比べる
        if not found: chart = "NG"; iss.append("C4_NO_CHART")
        else:
            chart = "OK"
            if len(found) < 3: chart = "HALF"; iss.append("C4_FEW_CHART")
            if bad: chart = "HALF"; iss.append("C4_NOT_SCATTER@" + ",".join(bad))
            if notitle: chart = "HALF"; iss.append("C4_NO_AXIS_TITLE@" + ",".join(notitle))
            sc = [nm for nm, d in (("横軸", xmax), ("縦軸", ymax)) if len(d) >= 2 and max(d.values()) - min(d.values()) > 1e-6]
            if sc: chart = "HALF"; iss.append("C4_SCALE@" + ",".join(sc))
        # ── 相関係数
        ok, hard, wrong = set(), set(), 0
        for row in ws.iter_rows():
            for c in row:
                v, val = c.value, wsv.cell(c.row, c.column).value
                if not num(val) or not -1 <= val <= 1: continue
                hit = [y for y in YEARS if abs(val - exp[y]) <= 0.005]
                isf = isinstance(v, str) and re.search(r"CORREL\(|PEARSON\(", v.upper())
                if isf and hit: ok.update(hit)
                elif isf: wrong += 1
                elif num(v) and hit and not (c.row in rows and 2 <= c.column <= 7): hard.update(hit)      # 元データ以外のセルに、相関係数の値が数値で入っている
        missing = [f"{y}年" for y in YEARS if y not in ok and y not in hard]
        if not ok and not hard and not wrong: corr = "NG"; iss.append("C4_NO_CORR")
        elif hard - ok: corr = "NG" if not ok else "HALF"; iss.append("C4_CORR_HARD")
        elif len(ok) == 3: corr = "OK"
        else:
            corr = "HALF" if ok else "NG"
            if wrong: iss.append("C4_CORR_WRONG@" + ",".join(missing))
            else: iss.append("C4_CORR_FEW@" + ",".join(missing))
        # ── 考察
        text, in_cell = essay_of(path, ws)
        n = len(nl(text))
        ess = ["C4_ESSAY_IN_CELL"] if in_cell else []
        if n < MIN_ESSAY_NONE: essay = "NG"; ess.append("C4_NO_ESSAY")
        else:
            lack = []
            if n < MIN_ESSAY: lack.append("C4_SHORT_ESSAY")
            if not re.search(r"相関|右上がり|右下がり", text): lack.append("C4_ESSAY_CORR")
            if sum(str(y) in text for y in YEARS) < 2: lack.append("C4_ESSAY_YEARS")
            causal = [s for s in re.split(r"[。\n]", text) if re.search(CAUSE, s) and not re.search(NEGATE, s)]
            if causal: lack.append("C4_CAUSAL")
            essay = "HALF" if lack else "OK"; ess += lack
        iss += ess                                             # コメントは散布図 → 相関係数 → 考察の順
        st = [chart, corr, essay]
        i = {0: 0, 1: 1, 2: 3, 3: 4}[st.count("NG")] + (1 if "HALF" in st else 0)
        J = {"OK": "できている", "HALF": "不十分", "NG": "できていない"}
        return ORDER[min(i, 4)], iss, {"散布図": J[chart], "相関係数": J[corr], "考察": J[essay], "考察文字数": n}
