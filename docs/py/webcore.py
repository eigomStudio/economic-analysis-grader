"""ブラウザ(Pyodide)用の採点の入口。grader.py の判定をそのまま使い、結果のExcelだけを簡素な形で作る。
入力は /tmp/in の下のxlsx、出力は指定のパスのxlsx。ネットワークには一切接続しない。"""
import re, sys, time, zipfile
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import grader as g
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as gl
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Meiryo"
HFILL = PatternFill("solid", fgColor="1F3864")
EDIT_FILL = PatternFill("solid", fgColor="BF8F00")
COLS = ["学生番号", "氏名", "課題", "採点結果", g.FINAL, g.COMMENT_COL, "できていない箇所", "ファイル名", "遅延"]
GRADES = ["S", "A", "B", "C", "D", "評価外", "未対応"]

def _hdr(ws, row, headers, fills=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row, j, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = (fills or {}).get(h, HFILL)
        c.alignment = Alignment(wrap_text=True, vertical="center")

def write_result(path, rows):
    wb = Workbook()
    ws = wb.active; ws.title = "採点結果"
    _hdr(ws, 1, COLS, {g.FINAL: EDIT_FILL, g.COMMENT_COL: EDIT_FILL})
    for r in rows:
        ws.append([r.get("学生番号", ""), r.get("氏名", ""), r.get("課題", ""), r["採点結果"], r["採点結果"],
                   g.make_comment([c for c in str(r.get("できていない箇所コード", "")).split(";") if c], r["採点結果"]),
                   r.get("できていない箇所", ""), r["ファイル"], r.get("遅延", "")])
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate([12, 14, 6, 10, 14, 55, 60, 40, 6], 1): ws.column_dimensions[gl(i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    last = max(ws.max_row, 2)
    dv = DataValidation(type="list", formula1='"S,A,B,C,D,評価外,未対応"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"E2:E{last}")
    for gname, col in (("D", "FFC7CE"), ("評価外", "FFC7CE"), ("未対応", "D9D9D9")):
        for L in "DE":
            ws.conditional_formatting.add(f"{L}2:{L}{last}", CellIsRule(operator="equal", formula=[f'"{gname}"'], fill=PatternFill("solid", bgColor=col)))
    # ── 集計(数式)。「できていない箇所」の列の文言を数える
    wsum = wb.create_sheet("集計")
    chapters = sorted({r["課題"] for r in rows if r.get("課題") not in ("", None)})
    wsum["A1"] = "採点結果の人数(最終評価)"; wsum["A1"].font = Font(name=FONT, bold=True, size=12)
    wsum["A2"] = f"採点プログラムのバージョン: {g.GRADER_VERSION}"; wsum["A2"].font = Font(name=FONT, size=9, color="888888")
    _hdr(wsum, 3, ["課題", "提出数(未対応を除く)"] + GRADES)
    for i, ch in enumerate(chapters, 4):
        wsum.cell(i, 1, ch)
        wsum.cell(i, 2, f"=COUNTIFS('採点結果'!$C:$C,$A{i},'採点結果'!$E:$E,\"<>未対応\")")
        for j in range(len(GRADES)):
            wsum.cell(i, 3 + j, f"=COUNTIFS('採点結果'!$C:$C,$A{i},'採点結果'!$E:$E,{gl(3 + j)}$3)")
    end1 = 3 + len(chapters)
    r0 = end1 + 3
    wsum.cell(r0 - 1, 1, "できていない箇所の人数").font = Font(name=FONT, bold=True, size=12)
    _hdr(wsum, r0, ["課題", "内容", "人数", "割合"])
    i = r0 + 1
    for code, label in g.LABELS.items():
        m = re.match(r"C(\d+)_", code)
        if not m or int(m.group(1)) not in chapters: continue
        wsum.cell(i, 1, int(m.group(1))); wsum.cell(i, 2, label)
        wsum.cell(i, 3, f"=COUNTIFS('採点結果'!$C:$C,$A{i},'採点結果'!$G:$G,\"*\"&$B{i}&\"*\")")
        wsum.cell(i, 4, f"=IFERROR(C{i}/INDEX($B$4:$B${end1},MATCH($A{i},$A$4:$A${end1},0)),0)")
        wsum.cell(i, 4).number_format = "0.0%"
        i += 1
    for row in wsum.iter_rows(min_row=4):
        for c in row:
            if not c.font.bold: c.font = Font(name=FONT, size=10)
    for k, w in enumerate([8, 60, 10, 8, 8, 8, 8, 8, 8], 1): wsum.column_dimensions[gl(k)].width = w
    wb.save(path)

def _zip_name(info):
    """ZIP内の日本語ファイル名の文字化けを直す(UTF-8の印がないZIPは、UTF-8またはShift_JIS系で復元を試みる)。"""
    n = info.filename
    if info.flag_bits & 0x800: return n
    for enc in ("utf-8", "cp932"):
        try: return n.encode("cp437").decode(enc)
        except Exception: pass
    return n

def expand_zips(in_dir):
    """Canvasの「全ての提出物をダウンロード」のZIPなど、ZIPに入った提出ファイルを取り出す。"""
    k = 0
    for z in sorted(Path(in_dir).rglob("*.zip")):
        with zipfile.ZipFile(z) as zf:
            for info in zf.infolist():
                name = Path(_zip_name(info)).name
                if info.is_dir() or "__MACOSX" in info.filename or name.startswith(("~$", ".")) or not name.lower().endswith(".xlsx"): continue
                k += 1; dest = Path(in_dir) / f"_zip{k}"; dest.mkdir(exist_ok=True); (dest / name).write_bytes(zf.read(info))
    return k

def run(in_dir, out_path, ref_path=None, chapter=None):
    t0 = time.time()
    g.load_plugins()
    expand_zips(in_dir)
    files = sorted((p for p in Path(in_dir).rglob("*.xlsx") if not p.name.startswith(("~$", "."))), key=lambda p: (p.name, str(p)))   # ファイル名順(ローカル版と同じ並び)
    refs = {}
    if ref_path:
        for n in g.CHECKERS: refs[n] = Path(ref_path)
    if not chapter:
        found = Counter(g.parse_name(p.stem)[2] or g.chapter_from_path(p.relative_to(in_dir)) for p in files)
        found.pop(None, None)
        chapter = found.most_common(1)[0][0] if found else None
        if not chapter and files:                    # ファイル名から章が分からない → 画面で章番号を尋ねる(採点はしない)
            return {"need_chapter": True, "count": len(files), "version": g.GRADER_VERSION}
    rows, log = [], []
    for p in files:
        sid, name, ch, fn_ok = g.parse_name(p.stem, g.chapter_from_path(p.relative_to(in_dir)) or chapter)
        base = {"学生番号": sid, "氏名": name, "課題": ch or "", "ファイル": p.name}
        if ch not in g.CHECKERS:
            rows.append(base | {"採点結果": "未対応", "できていない箇所": "", "できていない箇所コード": ""}); continue
        try:
            grade, iss, _ = g.CHECKERS[ch](p, refs.get(ch))
        except Exception as e:
            grade, iss = "評価外", []; log.append(f"{p.name}: 読み込めませんでした({str(e)[:60]})")
        if not fn_ok: iss = iss + ["FILENAME"]
        grade, iss, late = g.apply_file_late(p.relative_to(in_dir), grade, iss)   # ファイル名の「late」、フォルダ名の「遅延」 → 1段階下げる
        rows.append(base | {"採点結果": grade, "遅延": "あり" if late else "", "できていない箇所コード": ";".join(iss),
                            "できていない箇所": "、".join(g.label_of(c) for c in iss)})
    g.LABELS["FILENAME"] = "ファイル名から学生番号または氏名を読み取れない"
    for r in rows: r["できていない箇所"] = "、".join(g.label_of(c) for c in r["できていない箇所コード"].split(";") if c)
    write_result(out_path, rows)
    return {"version": g.GRADER_VERSION, "count": len(rows), "unsupported": sum(r["採点結果"] == "未対応" for r in rows), "chapter": chapter,
            "grades": dict(Counter(r["採点結果"] for r in rows)), "seconds": round(time.time() - t0, 2), "log": log}
